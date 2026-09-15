"""Multi-format document parser for file intelligence.

Supports: PDF (via pypdf), DOCX (via python-docx), XLSX/CSV (via openpyxl/csv),
JSON, plain text, and source code files.

All parsing is done in a fail-safe manner — if a library is missing,
the parser returns a graceful degradation message rather than crashing.
"""

import io
import os
import csv
import json
import logging
import mimetypes
from pathlib import Path
from typing import Optional, Tuple

logger = logging.getLogger(__name__)

# Maximum characters to extract per document (safety cap)
MAX_EXTRACT_CHARS = 12_000
MAX_ROWS_CSV = 200
MAX_SHEETS = 3


class ParsedDocument:
    """Structured result of document parsing."""

    __slots__ = ["filename", "content", "mime_type", "page_count", "metadata", "error"]

    def __init__(
        self,
        filename: str,
        content: str,
        mime_type: str = "application/octet-stream",
        page_count: int = 1,
        metadata: Optional[dict] = None,
        error: Optional[str] = None,
    ):
        self.filename = filename
        self.content = content
        self.mime_type = mime_type
        self.page_count = page_count
        self.metadata = metadata or {}
        self.error = error

    @property
    def is_success(self) -> bool:
        return self.error is None

    @property
    def truncated(self) -> bool:
        return len(self.content) >= MAX_EXTRACT_CHARS

    def to_context_string(self) -> str:
        """Format parsed document content for LLM context injection."""
        lines = [
            f"[DOCUMENT: {self.filename}]",
            f"[TYPE: {self.mime_type} | PAGES: {self.page_count}]",
        ]
        if self.metadata:
            meta_str = ", ".join(f"{k}={v}" for k, v in list(self.metadata.items())[:4])
            lines.append(f"[METADATA: {meta_str}]")
        if self.error:
            lines.append(f"[PARSE WARNING: {self.error}]")
        lines.append("")
        lines.append(self.content[:MAX_EXTRACT_CHARS])
        if self.truncated:
            lines.append(f"\n[... content truncated at {MAX_EXTRACT_CHARS} characters ...]")
        return "\n".join(lines)


class DocumentParser:
    """Multi-format document parser with graceful library degradation."""

    # ── Dispatch table ──────────────────────────────────────────────────────
    _EXTENSION_MAP = {
        ".pdf":   "_parse_pdf",
        ".docx":  "_parse_docx",
        ".doc":   "_parse_docx",
        ".xlsx":  "_parse_xlsx",
        ".xls":   "_parse_xlsx",
        ".csv":   "_parse_csv",
        ".json":  "_parse_json",
        ".jsonl": "_parse_json",
        ".txt":   "_parse_text",
        ".md":    "_parse_text",
        ".rst":   "_parse_text",
        ".yaml":  "_parse_text",
        ".yml":   "_parse_text",
        ".toml":  "_parse_text",
        ".ini":   "_parse_text",
        ".env":   "_parse_text",
        ".py":    "_parse_code",
        ".js":    "_parse_code",
        ".ts":    "_parse_code",
        ".jsx":   "_parse_code",
        ".tsx":   "_parse_code",
        ".java":  "_parse_code",
        ".go":    "_parse_code",
        ".rs":    "_parse_code",
        ".cpp":   "_parse_code",
        ".c":     "_parse_code",
        ".h":     "_parse_code",
        ".cs":    "_parse_code",
        ".swift": "_parse_code",
        ".kt":    "_parse_code",
        ".rb":    "_parse_code",
        ".php":   "_parse_code",
        ".sh":    "_parse_code",
        ".bash":  "_parse_code",
        ".sql":   "_parse_code",
    }

    def parse_bytes(self, filename: str, data: bytes) -> ParsedDocument:
        """Parse a document from raw bytes. Dispatches by file extension."""
        ext = Path(filename).suffix.lower()
        mime_type, _ = mimetypes.guess_type(filename)
        mime_type = mime_type or "application/octet-stream"

        method_name = self._EXTENSION_MAP.get(ext)
        if method_name:
            try:
                return getattr(self, method_name)(filename, data, mime_type)
            except Exception as exc:
                logger.warning("Parser %s failed for %s: %s", method_name, filename, exc)
                return ParsedDocument(
                    filename=filename,
                    content=self._try_text_fallback(data),
                    mime_type=mime_type,
                    error=f"Primary parser failed ({type(exc).__name__}), raw text fallback used.",
                )
        else:
            # Unknown extension — try UTF-8 text fallback
            return ParsedDocument(
                filename=filename,
                content=self._try_text_fallback(data),
                mime_type=mime_type,
                error="Unknown file type — raw text extraction used.",
            )

    def parse_file(self, filepath: str) -> ParsedDocument:
        """Parse a document from a local file path."""
        path = Path(filepath)
        if not path.exists():
            return ParsedDocument(
                filename=path.name,
                content="",
                error=f"File not found: {filepath}",
            )
        try:
            data = path.read_bytes()
            return self.parse_bytes(path.name, data)
        except PermissionError:
            return ParsedDocument(
                filename=path.name,
                content="",
                error=f"Permission denied reading: {filepath}",
            )
        except Exception as exc:
            return ParsedDocument(
                filename=path.name,
                content="",
                error=f"Read error: {type(exc).__name__}: {exc}",
            )

    # ── Parsers ─────────────────────────────────────────────────────────────

    def _parse_pdf(self, filename: str, data: bytes, mime_type: str) -> ParsedDocument:
        try:
            import pypdf
        except ImportError:
            return ParsedDocument(
                filename=filename,
                content=self._try_text_fallback(data),
                mime_type=mime_type,
                error="pypdf not installed. Run: pip install pypdf",
            )

        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append(text.strip())
            if sum(len(p) for p in pages) > MAX_EXTRACT_CHARS:
                break

        content = "\n\n---\n\n".join(pages)
        metadata = {}
        if reader.metadata:
            for k in ["Title", "Author", "Subject", "Creator"]:
                v = getattr(reader.metadata, k.lower(), None) or reader.metadata.get(f"/{k}")
                if v:
                    metadata[k] = str(v)[:100]

        return ParsedDocument(
            filename=filename,
            content=content[:MAX_EXTRACT_CHARS],
            mime_type=mime_type,
            page_count=len(reader.pages),
            metadata=metadata,
        )

    def _parse_docx(self, filename: str, data: bytes, mime_type: str) -> ParsedDocument:
        try:
            import docx as python_docx
        except ImportError:
            return ParsedDocument(
                filename=filename,
                content=self._try_text_fallback(data),
                mime_type=mime_type,
                error="python-docx not installed. Run: pip install python-docx",
            )

        doc = python_docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n".join(paragraphs)

        metadata = {}
        props = doc.core_properties
        if props.title:
            metadata["Title"] = str(props.title)[:100]
        if props.author:
            metadata["Author"] = str(props.author)[:100]

        return ParsedDocument(
            filename=filename,
            content=content[:MAX_EXTRACT_CHARS],
            mime_type=mime_type,
            page_count=1,
            metadata=metadata,
        )

    def _parse_xlsx(self, filename: str, data: bytes, mime_type: str) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError:
            return ParsedDocument(
                filename=filename,
                content=self._try_text_fallback(data),
                mime_type=mime_type,
                error="openpyxl not installed. Run: pip install openpyxl",
            )

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets_text = []
        for sheet_name in list(wb.sheetnames)[:MAX_SHEETS]:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS_CSV):
                row_str = "\t".join(str(c) if c is not None else "" for c in row)
                if row_str.strip():
                    rows.append(row_str)
            if rows:
                sheets_text.append(f"=== Sheet: {sheet_name} ===\n" + "\n".join(rows))

        content = "\n\n".join(sheets_text)
        return ParsedDocument(
            filename=filename,
            content=content[:MAX_EXTRACT_CHARS],
            mime_type=mime_type,
            page_count=len(wb.sheetnames),
            metadata={"sheets": ", ".join(wb.sheetnames[:MAX_SHEETS])},
        )

    def _parse_csv(self, filename: str, data: bytes, mime_type: str) -> ParsedDocument:
        try:
            text = data.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(text))
            rows = []
            for i, row in enumerate(reader):
                if i >= MAX_ROWS_CSV:
                    break
                rows.append(",".join(row))
            content = "\n".join(rows)
            return ParsedDocument(
                filename=filename,
                content=content[:MAX_EXTRACT_CHARS],
                mime_type=mime_type or "text/csv",
                page_count=1,
                metadata={"rows": str(min(i + 1, MAX_ROWS_CSV))},
            )
        except Exception as exc:
            return ParsedDocument(
                filename=filename,
                content=self._try_text_fallback(data),
                mime_type=mime_type,
                error=f"CSV parse error: {exc}",
            )

    def _parse_json(self, filename: str, data: bytes, mime_type: str) -> ParsedDocument:
        try:
            text = data.decode("utf-8", errors="replace")
            # Pretty-print to improve LLM readability
            parsed = json.loads(text)
            pretty = json.dumps(parsed, indent=2, ensure_ascii=False)
            return ParsedDocument(
                filename=filename,
                content=pretty[:MAX_EXTRACT_CHARS],
                mime_type=mime_type or "application/json",
                page_count=1,
            )
        except json.JSONDecodeError:
            # JSONL — parse line by line
            lines = data.decode("utf-8", errors="replace").splitlines()
            parsed_lines = []
            for line in lines[:50]:
                line = line.strip()
                if not line:
                    continue
                try:
                    obj = json.loads(line)
                    parsed_lines.append(json.dumps(obj, ensure_ascii=False))
                except Exception:
                    parsed_lines.append(line)
            return ParsedDocument(
                filename=filename,
                content="\n".join(parsed_lines)[:MAX_EXTRACT_CHARS],
                mime_type=mime_type or "application/jsonl",
                page_count=1,
                metadata={"lines": str(len(parsed_lines))},
            )

    def _parse_text(self, filename: str, data: bytes, mime_type: str) -> ParsedDocument:
        content = data.decode("utf-8", errors="replace")
        return ParsedDocument(
            filename=filename,
            content=content[:MAX_EXTRACT_CHARS],
            mime_type=mime_type or "text/plain",
            page_count=1,
        )

    def _parse_code(self, filename: str, data: bytes, mime_type: str) -> ParsedDocument:
        content = data.decode("utf-8", errors="replace")
        ext = Path(filename).suffix.lstrip(".")
        lang_label = {
            "py": "python", "js": "javascript", "ts": "typescript",
            "jsx": "jsx", "tsx": "tsx", "java": "java", "go": "go",
            "rs": "rust", "cpp": "c++", "c": "c", "cs": "csharp",
            "swift": "swift", "kt": "kotlin", "rb": "ruby", "php": "php",
            "sh": "bash", "bash": "bash", "sql": "sql",
        }.get(ext, ext)

        # Wrap in markdown code fence for LLM readability
        wrapped = f"```{lang_label}\n{content[:MAX_EXTRACT_CHARS]}\n```"
        return ParsedDocument(
            filename=filename,
            content=wrapped,
            mime_type=mime_type or "text/x-source",
            page_count=1,
            metadata={"language": lang_label},
        )

    def _try_text_fallback(self, data: bytes) -> str:
        """Best-effort UTF-8 decode of raw bytes."""
        try:
            return data.decode("utf-8", errors="replace")[:MAX_EXTRACT_CHARS]
        except Exception:
            return "[Binary content — cannot extract text]"


# ── Singleton ────────────────────────────────────────────────────────────────
_DOCUMENT_PARSER = DocumentParser()


def get_document_parser() -> DocumentParser:
    """Access the singleton DocumentParser."""
    return _DOCUMENT_PARSER
